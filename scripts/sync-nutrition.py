#!/usr/bin/env python3
"""Collect complete official K-FIND workbooks, preserving all source columns.

Requires openpyxl. No API key is required. The provider's download usage form is
submitted as individual / app DB development, without personal contact details.
Run --source-dir PATH to import previously downloaded official workbooks.
"""
from __future__ import annotations
import argparse
from datetime import date, datetime, timezone
import gzip
import hashlib
import html
from http.cookiejar import CookieJar
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
import time
from urllib.parse import urlencode
from urllib.request import build_opener, HTTPCookieProcessor, Request
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASE = 'https://various.foodsafetykorea.go.kr/nutrient/'
SOURCE_URL = BASE + 'general/down/historyList.do'
GROUPS = {'01': ('processed-foods', '가공식품 영양 DB'), '02': ('foods', '음식 영양 DB'), '12': ('supplements', '건강기능식품 영양 DB')}
BATCH_SIZE = 1000

def now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')

def clean(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return value if value is not None else ''

def encoded(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False).encode('utf-8')

def gzip_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    blob = gzip.compress(encoded(value), compresslevel=9, mtime=0)
    path.write_bytes(blob)
    return {'bytes': len(blob), 'sha256': hashlib.sha256(blob).hexdigest()}

def listing_versions(listing):
    versions = {}
    for row in re.findall(r'<tr[^>]*>(.*?)</tr>', listing, re.S):
        match = re.search(r"fnObj.fnDetail\('([0-9]+)', 'F', '', '([0-9]+)'", row)
        if match and match[2] in GROUPS:
            dates = re.findall(r'\d{4}-\d{2}-\d{2}', row)
            if dates:
                versions[match[2]] = dates[0]
    return versions

def fetch_workbooks(cache):
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    opener.addheaders = [('User-Agent', 'foodhannun/1.1'), ('Referer', SOURCE_URL)]
    def get(path, data=None):
        for attempt in range(3):
            try:
                return opener.open(Request(BASE + path, data=urlencode(data).encode() if data else None), timeout=150)
            except OSError:
                if attempt == 2:
                    raise RuntimeError('공식 다운로드에 연결하지 못했습니다.') from None
                time.sleep(2)
    listing = get('general/down/historyList.do').read().decode('utf-8')
    cache.mkdir(parents=True, exist_ok=True)
    (cache / 'listing.html').write_text(listing, encoding='utf-8')
    versions = {}
    for row in re.findall(r'<tr[^>]*>(.*?)</tr>', listing, re.S):
        match = re.search(r"fnObj.fnDetail\('([0-9]+)', 'F', '', '([0-9]+)'", row)
        if not match or match[2] not in GROUPS:
            continue
        seq, group = match.groups()
        dates = re.findall(r'\d{4}-\d{2}-\d{2}', row)
        version = dates[0] if dates else seq
        versions[group] = version
        dest = cache / f'{group}.xlsx'
        marker = cache / f'{group}.version'
        if dest.exists() and marker.exists() and marker.read_text() == version:
            continue
        popup = get('export/down/filePop.do?' + urlencode({'searchDbSeq': seq, 'searchDbGubun': group, 'searchOpenApi': 'Y'})).read().decode('utf-8')
        key = re.search(r'vmFileKey:"([^"]+)"', popup)
        token = re.search(r'name="pageToken"[^>]*value="([^"]+)"', popup)
        if not key or not token:
            raise RuntimeError('공식 다운로드 양식이 변경되었습니다.')
        registration = get('export/down/cudAction.do', {'pageToken': token[1], 'dbGubun': group, 'sosok': '음식한눈 개인개발', 'depart': '개발', 'agencyType': '01', 'usePurpose': '02', 'remark': ''}).read().decode('utf-8')
        if 'fnAfterCudAction' not in registration:
            raise RuntimeError('공식 다운로드 활용정보 등록 실패')
        partial = dest.with_suffix('.part')
        with get('multi/file/download.do?' + urlencode({'key': key[1]})) as response, partial.open('wb') as output:
            total = 0
            while chunk := response.read(1024 * 1024):
                total += len(chunk)
                if total > 300 * 1024 * 1024:
                    raise RuntimeError('원본 다운로드 크기 제한 초과')
                output.write(chunk)
        with partial.open('rb') as file:
            magic = file.read(2)
        if magic != b'PK':
            raise RuntimeError('원본 다운로드가 XLSX 파일이 아닙니다.')
        partial.replace(dest)
        marker.write_text(version)
        print(f'{GROUPS[group][1]} 원본 {total:,} bytes 다운로드', flush=True)
    if set(versions) != set(GROUPS):
        raise RuntimeError('공식 전체 DB 3종을 모두 발견하지 못했습니다.')
    return versions

def import_workbook(source, group, output, version=None):
    slug, title = GROUPS[group]
    workbook = load_workbook(source, read_only=True, data_only=True)
    source_hash = hashlib.file_digest(source.open('rb'), 'sha256').hexdigest()
    index, records, ids = [], [], set()
    chunks, total, expected = [], 0, 0
    columns = None
    date_values = set()
    def flush():
        if not records:
            return
        name = f'{slug}/part-{len(chunks):04d}.json.gz'
        stats = gzip_json(output / name, {'columns': columns, 'rows': records})
        chunks.append({'path': name, 'count': len(records), **stats})
        records.clear()
    for sheet in workbook:
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or '').strip() for value in next(rows)]
        if '식품코드' not in headers or '식품명' not in headers:
            raise ValueError(f'{title}: 필수 열이 없습니다.')
        if columns is not None and columns != headers:
            raise ValueError(f'{title}: 시트 간 열이 다릅니다.')
        columns = headers
        positions = {name: i for i, name in enumerate(headers)}
        code_pos, name_pos = positions['식품코드'], positions['식품명']
        category_pos = positions.get('식품대분류명')
        maker_pos = positions.get('제조사명')
        basis_pos = positions.get('영양성분함량기준량', positions.get('영양성분제공단위량'))
        date_pos = positions.get('데이터기준일자')
        expected += max(0, (sheet.max_row or 1) - 1)
        for row in rows:
            if not any(value not in (None, '') for value in row):
                continue
            code, name = str(row[code_pos] or '').strip(), str(row[name_pos] or '').strip()
            if not code or not name or code in ids:
                raise ValueError(f'{title}: 누락 또는 중복 식품코드')
            ids.add(code)
            # Preserve every nonempty source cell. Column numbers keep shard sizes small.
            sparse = {str(i): clean(value) for i, value in enumerate(row) if value not in (None, '')}
            index.append([code, name, row[category_pos] if category_pos is not None else '', row[maker_pos] if maker_pos is not None else '', len(chunks), len(records), row[basis_pos] if basis_pos is not None else ''])
            records.append(sparse)
            if date_pos is not None and row[date_pos]:
                date_values.add(str(clean(row[date_pos])))
            total += 1
            if len(records) >= BATCH_SIZE:
                flush()
            if total % 25000 == 0:
                print(f'{title}: {total:,}건 변환', flush=True)
    flush()
    workbook.close()
    if total != expected or not total:
        raise ValueError(f'{title}: 원본 {expected:,}행 / 변환 {total:,}건 불일치')
    index_path = f'{slug}/index.json.gz'
    index_stats = gzip_json(output / index_path, {'columns': ['id', 'name', 'category', 'manufacturer', 'chunk', 'row', 'basis'], 'rows': index})
    meta = {'id': slug, 'name': title, 'status': 'complete', 'count': total, 'sourceTotal': expected, 'uniqueIds': len(ids), 'columnCount': len(columns), 'sourceUrl': SOURCE_URL, 'version': version or (max(date_values) if date_values else None), 'fetchedAt': now(), 'sourceSha256': source_hash, 'sourceBytes': source.stat().st_size, 'sourceAttribution': '식품영양성분 데이터베이스 (K-FCDB), 식품의약품안전처', 'index': {'path': index_path, **index_stats}, 'chunks': chunks}
    print(f'{title}: {total:,}/{expected:,}건 완료, {len(columns)}개 원본 열 보존', flush=True)
    return meta

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source-dir', type=Path)
    parser.add_argument('--cache-dir', type=Path, default=ROOT / '.cache/nutrition')
    args = parser.parse_args()
    source_dir = args.source_dir or args.cache_dir
    versions = fetch_workbooks(source_dir) if not args.source_dir else listing_versions((source_dir / 'listing.html').read_text(encoding='utf-8')) if (source_dir / 'listing.html').exists() else {}
    target = ROOT / 'data/nutrition'
    target.parent.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(dir=target.parent, prefix='.nutrition-stage-') as staging:
        output = Path(staging)
        datasets = [import_workbook(source_dir / f'{group}.xlsx', group, output, versions.get(group)) for group in ['02', '12', '01']]
        manifest = {'schemaVersion': 1, 'generatedAt': now(), 'count': sum(item['count'] for item in datasets), 'sourceUrl': SOURCE_URL, 'datasets': datasets}
        (output / 'manifest.json').write_bytes(encoded(manifest))
        # Do not replace the last good generation until all workbooks validate.
        previous = target.with_name('.nutrition-previous')
        if previous.exists():
            shutil.rmtree(previous)
        if target.exists():
            target.rename(previous)
        try:
            output.rename(target)
        except OSError:
            if previous.exists():
                previous.rename(target)
            raise
        if previous.exists():
            shutil.rmtree(previous)
    print(f"공식 전체 영양 DB {manifest['count']:,}건 저장 완료", flush=True)

if __name__ == '__main__':
    main()
